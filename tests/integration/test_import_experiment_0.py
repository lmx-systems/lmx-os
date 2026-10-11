"""DE-3: the Experiment 0 import, on a workbook built here.

The real workbook carries customer data and stays on the shared drive, so these
build one with the same tabs and columns and invented businesses. What they pin
down is the shape of the import - which rows become which records, and that every
one grades reference - and that the records cannot leak into anything that runs
the business: the labelling queue, the nightly classifier, the scorecard, signup.
"""
import uuid
from datetime import datetime, timedelta, timezone

import openpyxl
import pytest
from sqlalchemy import func, select

from app.api.public_routes import client_signup
from app.identity.node_class import classify_unlabelled_locations, is_a_served_dock
from app.models.client import Client
from app.models.handoff_point import HandoffPoint
from app.models.hub import Hub
from app.models.location import Location
from app.models.site_address_share import SiteAddressShare
from app.models.stop import Stop
from app.reporting.operations import _hold_window_flag_rate
from app.schemas.signup import ClientSignupBody
from scripts.import_experiment_0 import (
    SOURCE_ROUTE_EXPORT,
    SOURCE_SITE_COUNTS,
    import_workbook,
    parse_workbook,
)

pytestmark = pytest.mark.integration

SITES = [
    "Site ID", "Source", "Account no.", "Business name", "Address", "Town", "ZIP",
    "Kind of site (from name, to confirm)", "Pickup or drop-off", "Shares address with",
    "Deliveries in data", "Tier today", "Notes from data",
]
DELIVERIES = [
    "Date", "Driver", "Manifest", "Account no.", "Business name", "Invoice", "Revenue ($)",
    "Pieces", "Weight", "Arrived", "Delivered", "Departed", "Stop time", "Stop time 10 s or less?",
]


def _workbook(path):
    book = openpyxl.Workbook()
    book.active.title = "Read me"
    book.active.append(["Experiment 0, invented for a test"])
    sites = book.create_sheet("Sites, any name")
    sites.append(SITES)
    sites.append(["T-1", "Source A", "100/0", "Elm Street Garage", "12 Elm Street", "Teaneck", "07666",
                  "Independent repair or tire shop", "Drop-off", None, 2, "Reference", None])
    # Two businesses at one address: one location, recorded as shared.
    sites.append(["T-2", "Source A", "200/0", "Cedar Body Works", "40 Cedar Lane", "Teaneck", "07666",
                  "Body shop", "Drop-off", "Cedar Tire Co", 1, "Reference", None])
    sites.append(["T-3", "Source A", "201/0", "Cedar Tire Co", "40 Cedar Lane", "Teaneck", "07666",
                  "Unclear from name", "Drop-off", "Cedar Body Works", 1, "Reference", None])
    # No address: its records have no site.
    sites.append(["T-4", "Source A", "300/0", "Somebody Auto", None, None, None,
                  "Independent repair or tire shop", "Drop-off", None, 1, "Reference", "No address"])
    sites.append(["T-5", "Source A", "900/0", "Parts Supplier", "1 Depot Road", "Bogota", "07603",
                  "Parts store or distributor counter", "Pickup likely (supplier)", None, 1, "Reference", None])
    # The other source: deliveries only as a count.
    sites.append(["T-6", "Source B", None, "Maple Dealer", "7 Maple Ave", "Paramus", "07652",
                  "Dealership service or parts counter", "Drop-off", None, 3, "Bronze at best", None])
    rows = book.create_sheet("Deliveries, any name")
    rows.append(DELIVERIES)
    for account, arrived in (("900/0", "8:05 AM"), ("100/0", "8:40 AM"), ("200/0", "9:10 AM"),
                             ("201/0", "9:10 AM"), ("100/0", "1:15 PM"), ("300/0", "2:00 PM")):
        rows.append(["Jan 02 2026", "A Real Name", "Manifest #1", account, "x", "N/A", 0, 2, 0,
                     arrived, arrived, arrived, "00:00:03", "Yes"])
    book.save(path)
    return path


async def _imported(db_session, tmp_path):
    report = await import_workbook(db_session, parse_workbook(_workbook(tmp_path / "exp0.xlsx")))
    await db_session.commit()
    return report


async def test_the_counts_reconcile_and_every_record_grades_reference(db_session, tmp_path):
    report = await _imported(db_session, tmp_path)

    assert (report.sites, report.dropoff_sites, report.pickup_sites) == (6, 5, 1)
    assert report.dropoff_sites_with_address == 4
    # T-2 and T-3 share an address, and T-4 has none.
    assert report.locations == 4
    assert (report.delivery_rows, report.delivery_rows_weighed, report.counted_deliveries) == (6, 0, 3)
    assert report.stops[(SOURCE_ROUTE_EXPORT, "pickup")] == 1
    assert report.stops[(SOURCE_ROUTE_EXPORT, "dropoff")] == 5
    assert report.stops[(SOURCE_SITE_COUNTS, "dropoff")] == 3
    # Records are drop-off stops, and none has a captured handoff point, an item,
    # an order or a label.
    assert dict(report.tiers) == {"reference": 8}


async def test_sites_carry_region_kind_and_the_shared_address(db_session, tmp_path):
    await _imported(db_session, tmp_path)

    locations = (await db_session.scalars(select(Location))).all()
    assert {loc.region for loc in locations} == {"NJ"}
    assert {loc.site_source for loc in locations} == {"import"}
    by_address = {loc.address.split(",")[0]: loc for loc in locations}
    assert by_address["12 Elm Street"].kind_of_site == "independent_repair_or_tire"
    assert by_address["12 Elm Street"].shares_address is None  # nobody checked
    # Two businesses resolved to one location: shared, and no pair with itself.
    assert by_address["40 Cedar Lane"].shares_address is True
    assert by_address["40 Cedar Lane"].kind_of_site == "body_shop"
    assert await db_session.scalar(select(func.count()).select_from(SiteAddressShare)) == 0
    # One placeholder door per location, nothing known about it.
    points = (await db_session.scalars(select(HandoffPoint))).all()
    assert len(points) == len(locations)
    assert all(p.handoff_type is None and p.source == "import" for p in points)


async def test_no_driver_name_is_stored_and_times_are_new_jersey_clock(db_session, tmp_path):
    await _imported(db_session, tmp_path)

    stops = (await db_session.scalars(select(Stop).where(Stop.record_source == SOURCE_ROUTE_EXPORT))).all()
    assert {s.collected_by for s in stops} == {"driver 1"}
    assert min(s.arrived_at for s in stops) == datetime(2026, 1, 2, 13, 5, tzinfo=timezone.utc)
    # The record with no site has no handoff point; the others do.
    assert sum(s.handoff_point_id is None for s in stops) == 1


async def test_running_it_twice_is_refused(db_session, tmp_path):
    await _imported(db_session, tmp_path)

    with pytest.raises(SystemExit, match="already been imported"):
        await import_workbook(db_session, parse_workbook(tmp_path / "exp0.xlsx"))


async def test_imported_sites_stay_out_of_the_labelling_queue_and_the_classifier(db_session, tmp_path):
    await _imported(db_session, tmp_path)

    served = await db_session.scalar(select(func.count()).select_from(Location).where(is_a_served_dock()))
    assert served == 0
    assert (await classify_unlabelled_locations(db_session))["considered"] == 0


async def test_imported_deliveries_do_not_dilute_the_held_wrong_rate(db_session, tmp_path):
    await _imported(db_session, tmp_path)

    rate = await _hold_window_flag_rate(db_session, datetime(2025, 1, 1, tzinfo=timezone.utc))
    assert rate.not_measured == "no completed deliveries in the window"


async def test_signup_never_lands_on_the_import_hub(db_session, real_redis_client, published_terms, tmp_path):
    # On a fresh database the import can run before the first real hub exists,
    # which makes its hub the oldest.
    await _imported(db_session, tmp_path)
    real = Hub(id=uuid.uuid4(), name="Austin Hub", lat=30.267, lng=-97.743,
               created_at=datetime.now(timezone.utc) + timedelta(seconds=1))
    db_session.add(real)
    await db_session.commit()

    body = ClientSignupBody(
        company_name="Midtown Auto Parts", contact_name="Jordan Rivera",
        contact_email=f"jordan-{uuid.uuid4().hex[:8]}@example.com", contact_phone="+15125550142",
        service_area="Austin metro", password="a-long-enough-password", terms_version="v1",
        accepted_terms=True,
    )
    request = type("R", (), {"client": type("C", (), {"host": "203.0.113.9"})()})()
    await client_signup(body, request, session=db_session)

    client = await db_session.scalar(select(Client).where(Client.name == "Midtown Auto Parts"))
    assert client.hub_id == real.id
